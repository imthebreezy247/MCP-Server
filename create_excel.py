import json
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

# Read the JSON data
with open('all_clients_quick.json', 'r') as f:
    clients = json.load(f)

# Create workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Daniel Berman Clients'

# Headers matching user's format
headers = ['Name', 'Policy Number', 'Phone', 'Email', 'Monthly Premium', 'Application Date',
           'Paid to Date', 'Gender', 'DOB', 'Age', 'Address', 'SSN', 'Date Received', 'Notes']

# Style headers
header_font = Font(bold=True, size=11)
header_fill = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
header_alignment = Alignment(horizontal='center', vertical='center')

# Write headers
for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment

# Write data
for row_num, client in enumerate(clients, 2):
    ws.cell(row=row_num, column=1, value=client.get('Name', ''))
    ws.cell(row=row_num, column=2, value=client.get('Policy Number', ''))
    ws.cell(row=row_num, column=3, value=client.get('Phone', ''))
    ws.cell(row=row_num, column=4, value=client.get('Email', ''))
    ws.cell(row=row_num, column=5, value=client.get('Premium', ''))
    ws.cell(row=row_num, column=6, value=client.get('Date', ''))
    ws.cell(row=row_num, column=7, value='')  # Paid to Date
    ws.cell(row=row_num, column=8, value='')  # Gender
    ws.cell(row=row_num, column=9, value='')  # DOB
    ws.cell(row=row_num, column=10, value='') # Age
    ws.cell(row=row_num, column=11, value='') # Address
    ws.cell(row=row_num, column=12, value='') # SSN
    ws.cell(row=row_num, column=13, value=client.get('Date', ''))  # Date Received
    ws.cell(row=row_num, column=14, value='')  # Notes

# Auto-size columns
for col in ws.columns:
    max_length = 0
    column = col[0].column_letter
    for cell in col:
        if cell.value:
            max_length = max(max_length, len(str(cell.value)))
    ws.column_dimensions[column].width = min(max_length + 2, 50)

# Save
wb.save('Daniel_Berman_Clients_All.xlsx')
print('Excel created: Daniel_Berman_Clients_All.xlsx')
print(f'Total clients: {len(clients)}')

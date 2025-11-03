#!/usr/bin/env python3
"""
Extract ALL ACA client deals from Daniel Berman emails
Export to Excel format for easy viewing
"""

import json
import re
import base64
from datetime import datetime
from pathlib import Path
import sys

# Google API imports
from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError

# Excel export
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    print("Warning: openpyxl not available. Install with: pip install openpyxl")

# Gmail API scope
SCOPES = ['https://www.googleapis.com/auth/gmail.readonly']

# Paths
SCRIPT_DIR = Path(__file__).parent.absolute()
TOKEN_PATH = SCRIPT_DIR / 'token.json'
CREDENTIALS_PATH = SCRIPT_DIR / 'credentials.json'


def proper_case_name(name):
    """Convert name to proper case"""
    if not name:
        return ""
    words = name.split()
    return ' '.join(word.capitalize() for word in words)


def extract_body(payload):
    """Extract email body from message payload"""
    body = ''

    if 'parts' in payload:
        for part in payload['parts']:
            if part['mimeType'] == 'text/plain' or part['mimeType'] == 'text/html':
                if 'data' in part['body']:
                    body += base64.urlsafe_b64decode(part['body']['data']).decode('utf-8', errors='ignore')
            elif 'parts' in part:
                body += extract_body(part)
    elif 'body' in payload and 'data' in payload['body']:
        body = base64.urlsafe_b64decode(payload['body']['data']).decode('utf-8', errors='ignore')

    return body


def extract_client_info(email_data):
    """Extract client information from email content"""

    body = email_data.get('body', '')
    snippet = email_data.get('snippet', '')
    subject = email_data.get('subject', '')
    content = body + '\n' + snippet

    client_info = {
        "Name": "",
        "Policy Number": "",
        "Phone": "",
        "Email": "",
        "Monthly Premium": "",
        "Application Date": "",
        "Paid to Date": "",
        "Gender": "",
        "DOB": "",
        "Age": "",
        "Address": "",
        "SSN": "",
        "Date Received": email_data.get('date', ''),
        "Notes": ""
    }

    # Extract Name from subject (most reliable for ACA emails)
    subject_name_match = re.search(r'^([A-Za-z\s]+?)\s+ACA', subject, re.IGNORECASE)
    if subject_name_match:
        name = subject_name_match.group(1).strip()
        client_info["Name"] = proper_case_name(name)

    # Extract Policy Number (52Z or 52Y followed by 7 digits, or other patterns)
    policy_patterns = [
        r'(52[ZY]\d{7})',
        r'(02[A-Z]\d{7})'  # Alternative pattern seen in data
    ]
    for pattern in policy_patterns:
        policy_match = re.search(pattern, content)
        if policy_match:
            client_info["Policy Number"] = policy_match.group(0)
            break

    # Extract Phone - comprehensive patterns
    phone_text = ""
    phone_patterns = [
        r'(?:Phone|Tel|Mobile|Cell|Ph)[:\s]*\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}',
        r'\(\d{3}\)\s*\d{3}[-.\s]?\d{4}',
        r'\d{3}[-.\s]\d{3}[-.\s]\d{4}',
        r'\b\d{10}\b'
    ]

    for pattern in phone_patterns:
        match = re.search(pattern, content, re.IGNORECASE)
        if match:
            phone_text = match.group(0)
            # Extract just the digits
            phone_digits = re.sub(r'\D', '', phone_text)
            if len(phone_digits) == 10:
                # Format as (XXX) XXX-XXXX
                client_info["Phone"] = f"({phone_digits[:3]}) {phone_digits[3:6]}-{phone_digits[6:]}"
                break

    # If phone has XXXX (partial), keep it as is
    if 'XXXX' in content or 'XXX' in content:
        partial_match = re.search(r'\(?\d{3}\)?[-.\s]?\d{3}[-.\s]X{3,4}', content, re.IGNORECASE)
        if partial_match:
            client_info["Phone"] = partial_match.group(0)

    # Extract Email address
    email_matches = re.findall(r'\b([a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,})\b', content)
    for email_addr in email_matches:
        if 'danielberman' not in email_addr.lower() and 'noreply' not in email_addr.lower():
            client_info["Email"] = email_addr.upper()
            break

    # Extract Monthly Premium
    premium_patterns = [
        r'(?:Monthly\s+Premium|Premium|Payment)[:\s]*\$\s*([\d,]+\.?\d{0,2})',
        r'\$\s*([\d,]+\.\d{2})',
    ]

    for pattern in premium_patterns:
        match = re.search(pattern, content, re.IGNORECASE)
        if match:
            premium = match.group(1).replace(',', '')
            if '.' not in premium:
                premium = premium + '.00'
            client_info["Monthly Premium"] = f"${premium}"
            break

    # Extract Application Date
    app_date_patterns = [
        r'(?:Application|App)[:\s]*(\d{1,2}/\d{1,2}/\d{4})',
    ]

    for pattern in app_date_patterns:
        match = re.search(pattern, content, re.IGNORECASE)
        if match:
            client_info["Application Date"] = match.group(1)
            break

    # Extract Paid to Date
    paid_patterns = [
        r'(?:Paid\s+[Tt]o)[:\s]*(\d{1,2}/\d{1,2}/\d{4})',
        r'Pending'
    ]

    for pattern in paid_patterns:
        match = re.search(pattern, content, re.IGNORECASE)
        if match:
            client_info["Paid to Date"] = match.group(0) if pattern == r'Pending' else match.group(1)
            break

    # Extract Gender
    gender_match = re.search(r'(?:Gender|Sex)[:\s]*(Male|Female|M|F)\b', content, re.IGNORECASE)
    if gender_match:
        gender = gender_match.group(1).upper()
        client_info["Gender"] = "Male" if gender in ['M', 'MALE'] else "Female"

    # Extract Date of Birth
    dob_patterns = [
        r'(?:DOB|Birth)[:\s]*(\d{1,2}/\d{1,2}/\d{4})',
    ]

    for pattern in dob_patterns:
        match = re.search(pattern, content, re.IGNORECASE)
        if match:
            client_info["DOB"] = match.group(1)
            break

    # Extract Age
    age_match = re.search(r'(?:Age)[:\s]*(\d{1,3})\b', content, re.IGNORECASE)
    if age_match:
        client_info["Age"] = age_match.group(1)

    # Extract Address - look for any address info
    address_patterns = [
        r'(\d+\s+[NSEW]?\s*[A-Z][A-Za-z\s.]+(?:Street|St|Ave|Avenue|Road|Rd|Drive|Dr|Lane|Ln|Boulevard|Blvd|Way|Court|Ct|STREET|ST|AVE|AVENUE|ROAD|RD|DRIVE|DR|LANE|LN|BOULEVARD|BLVD|WAY|COURT|CT)[^,\n]*)',
        r'(\d+\s+[A-Z]\s+[A-Z][A-Za-z\s]+)',  # Pattern like "1514 W CLINTON"
    ]

    for pattern in address_patterns:
        match = re.search(pattern, content, re.IGNORECASE)
        if match:
            address = match.group(1).strip()
            address = re.sub(r'\s+', ' ', address)
            client_info["Address"] = address.upper()
            break

    # Extract SSN
    ssn_match = re.search(r'(?:SSN)[:\s]*(\d{3}[-]?\d{2}[-]?\d{4})', content, re.IGNORECASE)
    if ssn_match:
        ssn = ssn_match.group(1)
        # Format as XXX-XX-XXXX
        ssn_digits = re.sub(r'\D', '', ssn)
        if len(ssn_digits) == 9:
            client_info["SSN"] = f"{ssn_digits[:3]}-{ssn_digits[3:5]}-{ssn_digits[5:]}"

    # Extract Notes - look for scheduling info, special instructions
    notes = []

    # Check for scheduling info
    time_patterns = [
        r'(\d{1,2}(?::\d{2})?\s*(?:am|pm|AM|PM))',
        r'(Available\s+[^\n]+)',
        r'(After\s+\d+[^\n]+)',
        r'(Tomorrow[^\n]*)',
        r'(anytime|ASAP|urgent)',
    ]

    for pattern in time_patterns:
        matches = re.finditer(pattern, content, re.IGNORECASE)
        for match in matches:
            note = match.group(1).strip()
            if note and note not in notes:
                notes.append(note)

    # Add any special notes from content
    special_patterns = [
        r'(Never\s+received[^\n]+)',
        r'(reschedule[^\n]*)',
        r'(Pending[^\n]*)',
    ]

    for pattern in special_patterns:
        match = re.search(pattern, content, re.IGNORECASE)
        if match:
            note = match.group(1).strip()
            if note and note not in notes:
                notes.append(note)

    client_info["Notes"] = "; ".join(notes[:3])  # Limit to 3 notes

    # Format Date Received to MM/DD/YYYY if possible
    date_received = client_info["Date Received"]
    try:
        # Try to parse various date formats
        from email.utils import parsedate_to_datetime
        dt = parsedate_to_datetime(date_received)
        client_info["Date Received"] = dt.strftime("%m/%d/%Y")
    except:
        pass

    return client_info


def authenticate():
    """Authenticate with Gmail API"""
    creds = None

    if TOKEN_PATH.exists():
        with open(TOKEN_PATH, 'r') as token:
            creds_data = json.load(token)
            creds = Credentials.from_authorized_user_info(creds_data, SCOPES)

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            print("Refreshing credentials...")
            creds.refresh(Request())
        else:
            if not CREDENTIALS_PATH.exists():
                raise FileNotFoundError(f"Missing {CREDENTIALS_PATH}")

            flow = InstalledAppFlow.from_client_secrets_file(str(CREDENTIALS_PATH), SCOPES)
            creds = flow.run_local_server(port=0)

        with open(TOKEN_PATH, 'w') as token:
            token.write(creds.to_json())

    return build('gmail', 'v1', credentials=creds)


def export_to_excel(client_data, output_file):
    """Export client data to Excel with formatting"""
    if not EXCEL_AVAILABLE:
        print("[WARNING] Excel export not available. Saving JSON only.")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ACA Clients"

    # Define headers
    headers = [
        "Name", "Policy Number", "Phone", "Email", "Monthly Premium",
        "Application Date", "Paid to Date", "Gender", "DOB", "Age",
        "Address", "SSN", "Date Received", "Notes"
    ]

    # Write headers with formatting
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Write data
    for row_num, client in enumerate(client_data, 2):
        ws.cell(row=row_num, column=1, value=client["Name"])
        ws.cell(row=row_num, column=2, value=client["Policy Number"])
        ws.cell(row=row_num, column=3, value=client["Phone"])
        ws.cell(row=row_num, column=4, value=client["Email"])
        ws.cell(row=row_num, column=5, value=client["Monthly Premium"])
        ws.cell(row=row_num, column=6, value=client["Application Date"])
        ws.cell(row=row_num, column=7, value=client["Paid to Date"])
        ws.cell(row=row_num, column=8, value=client["Gender"])
        ws.cell(row=row_num, column=9, value=client["DOB"])
        ws.cell(row=row_num, column=10, value=client["Age"])
        ws.cell(row=row_num, column=11, value=client["Address"])
        ws.cell(row=row_num, column=12, value=client["SSN"])
        ws.cell(row=row_num, column=13, value=client["Date Received"])
        ws.cell(row=row_num, column=14, value=client["Notes"])

    # Adjust column widths
    column_widths = {
        'A': 25,  # Name
        'B': 15,  # Policy Number
        'C': 18,  # Phone
        'D': 30,  # Email
        'E': 15,  # Monthly Premium
        'F': 15,  # Application Date
        'G': 15,  # Paid to Date
        'H': 10,  # Gender
        'I': 12,  # DOB
        'J': 8,   # Age
        'K': 40,  # Address
        'L': 15,  # SSN
        'M': 15,  # Date Received
        'N': 50   # Notes
    }

    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    # Freeze header row
    ws.freeze_panes = 'A2'

    # Save workbook
    wb.save(output_file)
    print(f"[OK] Excel file saved to: {output_file}")


def main():
    """Main function"""

    print("=" * 80)
    print("EXTRACTING ALL ACA CLIENT DEALS FROM GMAIL")
    print("=" * 80)

    # Authenticate
    print("\n[1/4] Authenticating...")
    try:
        service = authenticate()
        print("[OK] Authentication successful")
    except Exception as e:
        print(f"[ERROR] Authentication failed: {e}")
        return

    # Search for ALL emails
    print("\n[2/4] Searching for ALL ACA emails (excluding 'dead' label)...")
    query = 'from:danielberman.ushealth@gmail.com subject:ACA -label:dead'
    print(f"Query: {query}")

    try:
        all_messages = []
        page_token = None

        # Paginate through all results
        while True:
            results = service.users().messages().list(
                userId='me',
                q=query,
                maxResults=500,  # Max per page
                pageToken=page_token
            ).execute()

            messages = results.get('messages', [])
            all_messages.extend(messages)

            page_token = results.get('nextPageToken')
            print(f"  Retrieved {len(all_messages)} emails so far...")

            if not page_token:
                break

        total_emails = len(all_messages)
        print(f"[OK] Found {total_emails} total emails")

        if total_emails == 0:
            print("[WARNING] No emails found matching criteria")
            return

        # Process emails
        print(f"\n[3/4] Processing {total_emails} emails...")
        all_client_data = []

        for idx, msg in enumerate(all_messages, 1):
            if idx % 10 == 0 or idx == 1:
                print(f"  Processing email {idx}/{total_emails}...")

            try:
                # Get full email content
                msg_data = service.users().messages().get(
                    userId='me',
                    id=msg['id'],
                    format='full'
                ).execute()

                # Extract headers
                headers = {h['name']: h['value'] for h in msg_data['payload'].get('headers', [])}

                # Extract body
                body = extract_body(msg_data['payload'])

                email_data = {
                    'id': msg['id'],
                    'subject': headers.get('Subject', ''),
                    'date': headers.get('Date', ''),
                    'snippet': msg_data.get('snippet', ''),
                    'body': body
                }

                # Extract client information
                client_info = extract_client_info(email_data)
                all_client_data.append(client_info)

            except Exception as e:
                print(f"  [WARNING] Failed to process email {idx}: {e}")
                # Add placeholder
                all_client_data.append({
                    "Name": f"ERROR - Email {idx}",
                    "Policy Number": "",
                    "Phone": "",
                    "Email": "",
                    "Monthly Premium": "",
                    "Application Date": "",
                    "Paid to Date": "",
                    "Gender": "",
                    "DOB": "",
                    "Age": "",
                    "Address": "",
                    "SSN": "",
                    "Date Received": "",
                    "Notes": f"Failed to process: {str(e)}"
                })

        print(f"[OK] Processed {len(all_client_data)} emails")

        # Save outputs
        print("\n[4/4] Saving results...")

        # Save JSON
        json_file = SCRIPT_DIR / 'all_aca_deals.json'
        with open(json_file, 'w', encoding='utf-8') as f:
            json.dump(all_client_data, f, indent=2, ensure_ascii=False)
        print(f"[OK] JSON saved to: {json_file}")

        # Save Excel
        excel_file = SCRIPT_DIR / 'all_aca_deals.xlsx'
        export_to_excel(all_client_data, excel_file)

        print("\n" + "=" * 80)
        print("[SUCCESS] EXTRACTION COMPLETE!")
        print(f"Total records: {len(all_client_data)}")
        print(f"\nFiles created:")
        print(f"  - JSON: {json_file.absolute()}")
        print(f"  - Excel: {excel_file.absolute()}")
        print("=" * 80)

        # Show summary
        print("\nSummary:")
        with_policy = sum(1 for c in all_client_data if c["Policy Number"])
        with_phone = sum(1 for c in all_client_data if c["Phone"])
        with_email = sum(1 for c in all_client_data if c["Email"])
        print(f"  Records with Policy Number: {with_policy}")
        print(f"  Records with Phone: {with_phone}")
        print(f"  Records with Email: {with_email}")

    except HttpError as error:
        print(f"[ERROR] Gmail API error: {error}")
    except Exception as e:
        print(f"[ERROR] Error during processing: {e}")
        import traceback
        traceback.print_exc()


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n\n[INTERRUPTED] Stopped by user")
        sys.exit(1)
    except Exception as e:
        print(f"\n[FATAL ERROR] {e}")
        sys.exit(1)
